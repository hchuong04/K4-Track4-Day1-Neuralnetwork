"""results_table.py — Lưu kết quả huấn luyện ra JSON và điền tự động vào bảng experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl

FORMULA_COLS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    p = Path(results_dir)
    p.mkdir(parents=True, exist_ok=True)

    exp_id = result["cfg"].get("exp_id", "unnamed")
    out_path = p / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2)

    print(f"Đã lưu kết quả: {out_path}")
    return str(out_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    sm = result["summary"]
    exp_id = cfg.get("exp_id", "")

    hidden = cfg.get("hidden", (256, 128))
    hidden_str = "-".join(map(str, hidden))

    # Đảm bảo format các trường
    clip_norm_val = cfg.get("clip_norm")
    if clip_norm_val is None:
        clip_str = "none"
    else:
        clip_str = str(clip_norm_val)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce").upper(),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", 0.0),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": clip_str,
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": sm.get("step0_loss"),
        "best_val_loss": sm.get("best_val_loss"),
        "best_epoch": sm.get("best_epoch"),
        "final_train_loss": sm.get("final_train_loss"),
        "final_val_loss": sm.get("final_val_loss"),
        "val_acc": sm.get("val_acc"),
        "val_macro_f1": sm.get("val_macro_f1"),
        "time_per_epoch_s": sm.get("time_per_epoch_s"),
        "peak_mem_MB": sm.get("peak_mem_MB"),
        "diverged": "Có" if sm.get("diverged") else "Không",
        "eval_acc": eval_scores.get("accuracy") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("macro_f1") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes or cfg.get("notes", ""),
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               baseline_seeds: list[str] | None = None,
               group_comments: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments", "Seeds", "Summary" của mẫu rồi lưu thành out_path.

    Các bước:
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True
      2. ws = wb["Experiments"]; ghi dữ liệu vào các cột không phải công thức
      3. ws_seeds = wb["Seeds"]; ghi danh sách exp_id baseline vào cột A
      4. ws_summary = wb["Summary"]; ghi nhận xét ngắn vào cột H
      5. wb.save(out_path)
    """
    wb = openpyxl.load_workbook(template_path)

    # 1. Sheet Experiments
    ws = wb["Experiments"]
    header = [cell.value for cell in ws[1]]

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col_name in enumerate(header, start=1):
            if col_name and col_name in row and col_name not in FORMULA_COLS:
                val = row[col_name]
                ws.cell(row=r_idx, column=c_idx, value=val)

    # 2. Sheet Seeds (chỉ điền cột A các exp_id của baseline, các cột khác có công thức tự động)
    if "Seeds" in wb.sheetnames:
        ws_seeds = wb["Seeds"]
        if baseline_seeds is None:
            baseline_seeds = [r["exp_id"] for r in rows if r.get("group") == "baseline"]
        for idx, sid in enumerate(baseline_seeds[:5], start=2):
            ws_seeds.cell(row=idx, column=1, value=sid)

    # 3. Sheet Summary (điền nhận xét ngắn vào cột H: cột thứ 8)
    if "Summary" in wb.sheetnames and group_comments:
        ws_sum = wb["Summary"]
        for r in range(2, ws_sum.max_row + 1):
            grp = ws_sum.cell(row=r, column=1).value
            if grp in group_comments:
                ws_sum.cell(row=r, column=8, value=group_comments[grp])

    wb.save(out_path)
    print(f"Đã xuất bảng thành công: {out_path} ({len(rows)} thí nghiệm)")
